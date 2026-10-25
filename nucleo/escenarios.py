"""Escenarios sin copiar archivos: un modelo base más una capa de cambios guardada en JSON.

Un escenario dice qué proyectos van prendidos o apagados, estados y consignas puntuales, demandas (por ejemplo,
importadas del Excel de demanda proyectada), el multiplicador de demanda, opciones y las ediciones de geometría
hechas en el mapa. aplicar() arma el .inp resultante con el Editor, sin tocar nada más del archivo base.
"""
import json
import math
import os
import re
import time
from collections import OrderedDict, defaultdict

from . import convenciones, inp

VACIO = {'nombre': 'Sin cambios', 'base': '', 'notas': '', 'ediciones': [], 'proyectos': {}, 'estados': {},
         'consignas': {}, 'demandas': {}, 'multiplicador': None, 'opciones': []}


def nuevo(base, nombre='Sin cambios'):
    e = json.loads(json.dumps(VACIO))
    e['base'] = base
    e['nombre'] = nombre
    e['creado'] = time.strftime('%Y-%m-%d %H:%M')
    return e


def vacio(esc):
    return not any(esc.get(k) for k in ('ediciones', 'proyectos', 'estados', 'consignas', 'demandas', 'opciones')) \
        and esc.get('multiplicador') in (None, '')


# ------------------------------------------------------------------ grupos de proyecto
def grupos(m, conv=None):
    """Grupos para prender/apagar: horizontes (año según la etiqueta), proyectos (según el ID, si hay una regla en
    Ajustes) y las demás etiquetas con nombre (las solo numéricas suelen ser sectores y no se agrupan)."""
    conv = convenciones.normalizar(conv)
    por = OrderedDict()
    for l in m.tramos:
        t = m.tag('tramo', l)
        p = convenciones.proyecto(l, conv)
        if p:
            por.setdefault(('proyecto', p), []).append(l)
        if t:
            a = convenciones.horizonte(t, conv)
            if a:
                por.setdefault(('horizonte', str(a)), []).append(l)
            elif not re.match(r'^[\d.]+$', t):
                por.setdefault(('etiqueta', t), []).append(l)
    out = []
    for (tipo, nombre), ls in por.items():
        if tipo == 'etiqueta' and len(ls) < 2:
            continue
        cerr = sum(1 for l in ls if m.cerrado_al_inicio(l))
        horiz = sorted({str(convenciones.horizonte(m.tag('tramo', l), conv)) for l in ls
                        if convenciones.horizonte(m.tag('tramo', l), conv)})
        out.append({'clave': f'{tipo}:{nombre}', 'tipo': tipo, 'nombre': nombre, 'tramos': ls, 'n': len(ls),
                    'cerrados': cerr, 'horizontes': horiz,
                    'estado': 'apagado' if cerr == len(ls) else ('prendido' if cerr == 0 else 'mixto')})
    orden = {'horizonte': 0, 'proyecto': 1, 'etiqueta': 2}

    def clave(g):
        n = re.search(r'(\d+)', g['nombre'])
        return (orden[g['tipo']], int(n.group(1)) if n else 0, g['nombre'])
    return sorted(out, key=clave)


def estados_grupo(m, tramos, accion, abrir_bombas=True):
    """Estados para prender o apagar un grupo.

    Prender: tuberías y bombas abiertas, válvulas con su consigna; con abrir_bombas=False las bombas quedan como en
    el archivo (en muchas bases vienen cerradas y las prenden sus reglas; las que no tienen regla son de reserva).
    Apagar: se cierran los tramos que tocan la red existente y las bombas; si queda una parte sin tanque, se le deja
    abierta una sola entrada (la tubería de mayor diámetro hacia la red), que es lo que evita el error 110 sin que
    circule agua."""
    G = set(tramos)
    est = {}
    if accion == 'prendido':
        for l in G:
            t = m.tramos[l]['tipo']
            if t == 'pump' and not abrir_bombas:
                continue
            est[l] = 'OPEN' if t in ('pipe', 'pump') else ''
        return est
    nudos_g = defaultdict(list)
    fuera = set()
    for l, d in m.tramos.items():
        for n in (d['n1'], d['n2']):
            if l in G:
                nudos_g[n].append(l)
            else:
                fuera.add(n)
    frontera = {n for n in nudos_g if n in fuera}
    borde = {l for n in frontera for l in nudos_g[n]}
    for l in borde:
        est[l] = 'CLOSED'
    for l in G:
        if m.tramos[l]['tipo'] == 'pump':
            est[l] = 'CLOSED'
    # partes que quedan sin tanque: se les deja una entrada
    ady = defaultdict(list)
    for l in G:
        if est.get(l) == 'CLOSED' or (l not in est and m.cerrado_al_inicio(l)):
            continue
        d = m.tramos[l]
        ady[d['n1']].append(d['n2']); ady[d['n2']].append(d['n1'])
    visto = set()
    for n0 in nudos_g:
        if n0 in visto or n0 in frontera:
            continue
        comp, pila = set(), [n0]
        visto.add(n0)
        while pila:
            x = pila.pop(); comp.add(x)
            for y in ady[x]:
                if y not in visto and y not in frontera:
                    visto.add(y); pila.append(y)
        if any(m.nudos.get(x, {}).get('tipo') in ('tank', 'reservoir') for x in comp):
            continue
        cand = [l for l in borde if (m.tramos[l]['n1'] in comp) != (m.tramos[l]['n2'] in comp)
                and m.tramos[l]['tipo'] != 'pump']
        if not cand:
            continue
        cand.sort(key=lambda l: (m.tramos[l]['tipo'] != 'pipe', -m.tramos[l].get('diam', 0)))
        e = cand[0]
        est[e] = 'OPEN' if m.tramos[e]['tipo'] == 'pipe' else ''
    return est


# ------------------------------------------------------------------ aplicar
def aplicar(base, esc, conv=None):
    """(texto, avisos, editor): el .inp del escenario y lo que no se pudo aplicar."""
    e = inp.Editor(base, conv)
    avisos = []
    for k, op in enumerate(esc.get('ediciones', [])):
        try:
            e.aplicar(op)
        except (inp.ErrorEdicion, KeyError, ValueError, TypeError) as ex:
            avisos.append(f'Edición {k + 1} ({op.get("op")}): {ex}')
    if esc.get('proyectos'):
        gs = {g['clave']: g for g in grupos(base, conv)}
        # tuberías que un quinquenio no apaga aunque su etiqueta sea de después (las que cambia un .scn existen antes
        # de su año: ver nucleo/quinquenios.py)
        conservar = set((esc.get('quinquenio') or {}).get('conservar') or [])
        # en un quinquenio las bombas de lo que se prende quedan como en el archivo: las manejan sus reglas
        abrir_bombas = not esc.get('quinquenio')
        # primero los horizontes, después los proyectos: lo más específico manda
        orden = sorted(esc['proyectos'].items(), key=lambda kv: (0 if kv[0].startswith('horizonte') else 1, kv[0]))
        for clave, accion in orden:
            g = gs.get(clave)
            if not g or accion not in ('prendido', 'apagado'):
                continue
            tramos = [l for l in g['tramos'] if e.existe_tramo(l) and not (accion == 'apagado' and l in conservar)]
            for l, v in estados_grupo(base, tramos, accion, abrir_bombas).items():
                e.op_estado(l, v)
    for l, v in esc.get('estados', {}).items():
        try:
            e.op_estado(l, v)
        except inp.ErrorEdicion as ex:
            avisos.append(str(ex))
    for l, v in esc.get('consignas', {}).items():
        try:
            d = e.tramo(l)
            if d['tipo'] != 'valve':
                raise inp.ErrorEdicion(f'{l} no es una válvula')
            e.op_propiedad('tramo', l, 'consigna', v)
            ev = base.estado.get(l)
            if ev and ev[1].upper() not in ('OPEN', 'CLOSED') and l not in esc.get('estados', {}):
                e.op_estado(l, inp.fmt(float(v)))     # una consigna en [STATUS] manda sobre la de [VALVES]
        except (inp.ErrorEdicion, ValueError) as ex:
            avisos.append(str(ex))
    for n, v in esc.get('demandas', {}).items():
        try:
            q, pat = (v + [None])[:2] if isinstance(v, list) else (v, None)
            e.op_propiedad('nudo', n, 'dem', q)
            if pat:
                e.op_propiedad('nudo', n, 'pat', pat)
        except (inp.ErrorEdicion, ValueError) as ex:
            avisos.append(str(ex))
    if esc.get('multiplicador') not in (None, ''):
        e.op_multiplicador(esc['multiplicador'])
    for o in esc.get('opciones', []):
        e.op_opcion(o['seccion'], o['clave'], o['valor'])
    return e.texto(), avisos + e.avisos, e


# ------------------------------------------------------------------ Excel
def _norm_id(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def hojas_excel(ruta, filas=8):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        muestra = []
        n = 0
        for r in ws.iter_rows(values_only=True):
            n += 1
            if len(muestra) < filas:
                muestra.append(['' if v is None else (str(v)[:40]) for v in r[:30]])
            elif n > 20000:
                break
        out.append({'nombre': ws.title, 'filas': n, 'muestra': muestra})
    wb.close()
    return out


def leer_tabla(ruta, hoja, fila_encabezado, col_id, col_valor, col_patron=None):
    """Filas (id, valor, patrón) de una hoja; las columnas van numeradas desde 1, como en Excel (A = 1)."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[hoja]
    out, malos = [], 0
    for k, r in enumerate(ws.iter_rows(values_only=True), 1):
        if k <= fila_encabezado:
            continue
        r = list(r) + [None] * 40
        ident = _norm_id(r[col_id - 1])
        if not ident or ident.startswith(';') or ident.startswith('['):
            continue
        try:
            val = float(r[col_valor - 1])
        except (TypeError, ValueError):
            malos += 1
            continue
        pat = _norm_id(r[col_patron - 1]) if col_patron else None
        out.append((ident, val, pat or None))
    wb.close()
    return out, malos


def importar(esc, modelo, filas, destino):
    """Pone las filas leídas en el escenario. Solo entran los IDs que existen y son del tipo correcto."""
    ok, desconocidos, sin_patron = 0, [], set()
    for ident, val, pat in filas:
        if destino == 'demandas':
            if modelo.nudos.get(ident, {}).get('tipo') != 'junction':
                desconocidos.append(ident); continue
            if pat and pat not in modelo.patrones:      # p. ej. "Sin asignar aún": se deja el patrón del nudo
                sin_patron.add(pat); pat = None
            esc.setdefault('demandas', {})[ident] = [val, pat]
        elif destino == 'consignas':
            if modelo.tramos.get(ident, {}).get('tipo') != 'valve':
                desconocidos.append(ident); continue
            esc.setdefault('consignas', {})[ident] = val
        else:
            raise ValueError(f'Destino desconocido: {destino}')
        ok += 1
    return {'aplicados': ok, 'desconocidos': len(desconocidos), 'muestra_desconocidos': desconocidos[:15],
            'patrones_ignorados': sorted(sin_patron)}


# ------------------------------------------------------------------ guardar
def ruta_para(nombre, carpeta):
    limpio = re.sub(r'[^\w\- ]+', '_', nombre or 'escenario').strip() or 'escenario'
    return os.path.join(carpeta, limpio + '.json')


def guardar(esc, carpeta, ruta=None):
    ruta = ruta or ruta_para(esc.get('nombre'), carpeta)
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    esc['guardado'] = time.strftime('%Y-%m-%d %H:%M')
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(esc, fh, ensure_ascii=False, indent=1)
    return ruta


def listar(carpeta):
    out = []
    if not os.path.isdir(carpeta):
        return out
    for f in sorted(os.listdir(carpeta)):
        if f.endswith('.json'):
            try:
                with open(os.path.join(carpeta, f), encoding='utf-8') as fh:
                    e = json.load(fh)
                out.append({'archivo': os.path.join(carpeta, f), 'nombre': e.get('nombre', f[:-5]),
                            'base': e.get('base', ''), 'guardado': e.get('guardado', ''),
                            'cambios': sum(len(e.get(k) or []) for k in ('ediciones', 'proyectos', 'estados', 'consignas', 'demandas', 'opciones'))})
            except (OSError, ValueError):
                pass
    return out
