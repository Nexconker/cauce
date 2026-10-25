"""Armar un quinquenio con un clic: un escenario con las obras hasta ese año prendidas y las de después apagadas, los
escenarios de EPANET (.scn) que van desde un año y la demanda de ese año.

La configuración es una sola para todos los modelos (config.json, clave "quinquenio"), porque es del proyecto y no de
una versión del .inp:
- scn: [{ruta, desde}]. Las tuberías que nombra un .scn existen antes de su año (el .scn solo les cambia el diámetro
  o la rugosidad): por eso nunca se apagan aunque su etiqueta sea de un horizonte posterior (pasa con una línea
  existente que se etiquetó con el año en que cambia de diámetro).
- demanda: reparto por sector, como una distribución de demanda proyectada: una hoja con nudo, sector y
  ponderación (la parte del sector que lleva el nudo) y otra con la demanda de cada sector por año (columnas con el
  año en una fila de encabezado; si un sector aparece en varias filas, se suman). Demanda del nudo = ponderación ×
  demanda del sector en el año.
"""
import os
import re

from . import convenciones, escenarios, inp, scn

CLAVES_DEMANDA = ('ruta', 'hoja_nudos', 'fila_nudos', 'col_id', 'col_sector', 'col_peso', 'hoja_sectores',
                  'fila_anios', 'col_sector_s')


def normalizar(cfg):
    """La configuración con lo que falte vacío y los tipos limpios."""
    cfg = cfg or {}
    out = {'scn': [], 'demanda': None}
    for x in cfg.get('scn') or []:
        try:
            out['scn'].append({'ruta': str(x['ruta']), 'desde': int(x['desde'])})
        except (KeyError, TypeError, ValueError):
            pass
    out['scn'].sort(key=lambda x: (x['desde'], x['ruta'].lower()))
    d = cfg.get('demanda')
    if isinstance(d, dict) and all(d.get(k) not in (None, '') for k in CLAVES_DEMANDA):
        try:
            out['demanda'] = {'ruta': str(d['ruta']), 'hoja_nudos': str(d['hoja_nudos']), 'hoja_sectores': str(d['hoja_sectores']),
                              **{k: int(d[k]) for k in ('fila_nudos', 'col_id', 'col_sector', 'col_peso', 'fila_anios', 'col_sector_s')},
                              'col_patron': int(d['col_patron']) if d.get('col_patron') not in (None, '') else None}
        except (TypeError, ValueError):
            out['demanda'] = None
    return out


def anios(m, conv=None):
    """Los años de los horizontes del modelo (las etiquetas de sus tramos)."""
    conv = convenciones.normalizar(conv)
    return sorted({a for a in (convenciones.horizonte(m.tag('tramo', l), conv) for l in m.tramos) if a})


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _es_anio(v, anio):
    if isinstance(v, (int, float)):
        return float(v) == float(anio)
    return isinstance(v, str) and v.strip() == str(anio)


def demanda_del_anio(cfg, anio):
    """({nudo: (demanda, patrón o None)}, info). info dice si encontró la columna del año y qué no cuadró."""
    import openpyxl
    if not os.path.isfile(cfg['ruta']):
        raise ValueError(f'No encuentro el Excel de demanda: {cfg["ruta"]}')
    wb = openpyxl.load_workbook(cfg['ruta'], data_only=True, read_only=True)
    try:
        if cfg['hoja_sectores'] not in wb.sheetnames or cfg['hoja_nudos'] not in wb.sheetnames:
            raise ValueError('El Excel de demanda ya no tiene las hojas configuradas.')
        sectores, col_anio, disponibles = {}, None, []
        for k, r in enumerate(wb[cfg['hoja_sectores']].iter_rows(values_only=True), 1):
            r = list(r)
            if k == cfg['fila_anios']:
                # la primera columna con ese año a la derecha de la del sector (a veces hay otra tabla
                # más a la derecha con los mismos años)
                for j in range(cfg['col_sector_s'], len(r)):
                    v = r[j]
                    if isinstance(v, (int, float)) and 1900 < v < 2200 and float(v).is_integer():
                        disponibles.append(int(v))
                    if col_anio is None and _es_anio(v, anio):
                        col_anio = j
                continue
            if k < cfg['fila_anios'] or col_anio is None or len(r) <= col_anio:
                continue
            s = escenarios._norm_id(r[cfg['col_sector_s'] - 1])
            q = _num(r[col_anio])
            if s and q is not None:
                sectores[s] = sectores.get(s, 0.0) + q
        info = {'anio': anio, 'encontrado': col_anio is not None, 'anios': sorted(set(disponibles)),
                'sectores': len(sectores), 'sin_sector': [], 'filas': 0}
        if col_anio is None:
            return {}, info
        out = {}
        for k, r in enumerate(wb[cfg['hoja_nudos']].iter_rows(values_only=True), 1):
            if k <= cfg['fila_nudos']:
                continue
            r = list(r) + [None] * 40
            n = escenarios._norm_id(r[cfg['col_id'] - 1])
            if not n or n.startswith((';', '[')):
                continue
            s = escenarios._norm_id(r[cfg['col_sector'] - 1])
            w = _num(r[cfg['col_peso'] - 1])
            if w is None:
                continue
            info['filas'] += 1
            if s not in sectores:
                info['sin_sector'].append(n)
                continue
            pat = escenarios._norm_id(r[cfg['col_patron'] - 1]) if cfg.get('col_patron') else ''
            q0, p0 = out.get(n, (0.0, None))
            out[n] = (q0 + w * sectores[s], p0 or pat or None)
        return out, info
    finally:
        wb.close()


def armar(base, anio, cfg, conv=None, nombre=None):
    """(escenario, resumen) del quinquenio `anio` sobre el modelo base."""
    anio = int(anio)
    cfg = normalizar(cfg)
    esc = escenarios.nuevo(base.ruta, nombre or f'Quinquenio {anio}')
    resumen = {'anio': anio, 'prendidos': [], 'apagados': [], 'conservados': 0, 'scn': [], 'demanda': None, 'avisos': []}
    # .scn: los de su año para atrás se aplican; todas sus tuberías se conservan (existen antes de su año)
    conservar = set()
    for x in cfg['scn']:
        r = {'nombre': os.path.basename(x['ruta']), 'ruta': x['ruta'], 'desde': x['desde'], 'aplica': x['desde'] <= anio}
        try:
            datos = scn.leer_archivo(x['ruta'])
        except (OSError, ValueError) as ex:
            r['error'] = str(ex)
            resumen['scn'].append(r)
            resumen['avisos'].append(f'{r["nombre"]}: {ex}')
            continue
        tuberias = {f[0] for s in ('DIAMETERS', 'ROUGHNESS') for f in datos['secciones'].get(s, [])}
        conservar |= {t for t in tuberias if base.tramos.get(t, {}).get('tipo') == 'pipe'}
        r['tuberias'] = len(tuberias)
        if r['aplica']:
            secciones = datos['secciones']
            malos = inp.Editor(base, conv).revisar_scn(secciones)
            if malos:
                secciones = scn.quitar(secciones, malos)
                r['omitidos'] = len(malos)
            if secciones:
                esc['ediciones'].append({'op': 'scn', 'nombre': r['nombre'], 'ruta': x['ruta'], 'notas': datos['notas'],
                                         'secciones': secciones})
        resumen['scn'].append(r)
    # obras: los horizontes hasta el año prendidos, los de después apagados (con una entrada a cada parte sin tanque)
    for g in escenarios.grupos(base, conv):
        if g['tipo'] != 'horizonte':
            continue
        a = int(g['nombre'])
        accion = 'prendido' if a <= anio else 'apagado'
        esc['proyectos'][g['clave']] = accion
        n = len([l for l in g['tramos'] if not (accion == 'apagado' and l in conservar)])
        resumen['prendidos' if accion == 'prendido' else 'apagados'].append({'anio': a, 'tramos': n})
    conv_n = convenciones.normalizar(conv)
    # las que se habrían apagado sin el .scn
    resumen['conservados'] = sum(1 for l in conservar if (convenciones.horizonte(base.tag('tramo', l), conv_n) or 0) > anio)
    # demanda del año
    if cfg['demanda']:
        try:
            dem, info = demanda_del_anio(cfg['demanda'], anio)
        except (OSError, ValueError, KeyError) as ex:
            dem, info = {}, {'error': str(ex)}
            resumen['avisos'].append(f'Demanda: {ex}')
        if dem:
            r = escenarios.importar(esc, base, [(n, q, p) for n, (q, p) in dem.items()], 'demandas')
            antes = sum(base.nudos[n]['dem'] for n in esc['demandas'])
            despues = sum(v[0] for v in esc['demandas'].values())
            info.update(nudos=r['aplicados'], desconocidos=r['desconocidos'], muestra_desconocidos=r['muestra_desconocidos'],
                        patrones_ignorados=r['patrones_ignorados'], antes=antes, despues=despues,
                        total_modelo=sum(d['dem'] for d in base.nudos.values() if d['tipo'] == 'junction'))
        info['archivo'] = os.path.basename(cfg['demanda']['ruta'])
        sin = info.get('sin_sector') or []
        info['sin_sector_n'], info['sin_sector'] = len(sin), sin[:15]
        resumen['demanda'] = info
        if info.get('encontrado') is False:
            resumen['avisos'].append(f'El Excel de demanda no tiene la columna de {anio}'
                                     + (f' (tiene {", ".join(map(str, info["anios"]))})' if info.get('anios') else '') + ': la demanda queda como en el archivo.')
    esc['quinquenio'] = {'anio': anio, 'conservar': sorted(conservar),
                         'scn': [r['nombre'] for r in resumen['scn'] if r.get('aplica') and 'error' not in r],
                         'demanda': bool(esc['demandas'])}
    return esc, resumen


def nombre_valido(nombre):
    return re.sub(r'\s+', ' ', str(nombre or '')).strip()[:80]
